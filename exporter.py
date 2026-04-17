from __future__ import annotations

import os
from dataclasses import dataclass
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import (
    EXCEL_HEADER_COLOR,
    EXCEL_HEADER_FONT_COLOR,
    EXCEL_ROW_ALT_COLOR,
    EXCEL_ROW_BASE_COLOR,
    LEADS_SHEET_NAME,
    OUTPUT_COLUMNS,
    SOCIAL_SEARCH_COLUMNS,
    SOCIAL_SEARCH_SHEET_NAME,
    SUMMARY_SHEET_NAME,
    UPLOAD_READY_SHEET_NAME,
    UNIVERSAL_FINDER_COLUMNS,
    UNIVERSAL_FINDER_SHEET_NAME,
    VERSION,
)
from utils import deduplicate_leads, ensure_parent_dir, now_string

STATUS_OPTIONS = ["Pending", "Complete", "Decline"]
STATUS_COLORS = {
    "Pending": ("FFF4CC", "7A5C00"),
    "Complete": ("D9EAD3", "1E6B34"),
    "Decline": ("F4CCCC", "9C0006"),
}
SHEET_ACCENT_COLOR = "0E7490"
SUMMARY_LABEL_FILL = "DFF6FD"
SUMMARY_VALUE_FILL = "F8FDFF"
ENGAGEMENT_BAND_FILL = "E6FFFB"
ENGAGEMENT_BAND_TEXT = "0F172A"
THIN_BORDER = Border(
    left=Side(style="thin", color="D7E3F1"),
    right=Side(style="thin", color="D7E3F1"),
    top=Side(style="thin", color="D7E3F1"),
    bottom=Side(style="thin", color="D7E3F1"),
)


def _add_status_validation(worksheet, header_index: dict[str, int]) -> None:
    status_col = header_index.get("Lead Status")
    if not status_col:
        return
    formula = '"' + ",".join(STATUS_OPTIONS) + '"'
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.prompt = "Choose a lead review status."
    validation.promptTitle = "Lead Status"
    validation.error = "Pick a valid status from the list."
    validation.errorTitle = "Invalid Status"
    worksheet.add_data_validation(validation)
    validation.add(f"{worksheet.cell(row=3, column=status_col).column_letter}3:{worksheet.cell(row=max(worksheet.max_row, 5001), column=status_col).column_letter}{max(worksheet.max_row, 5001)}")

    status_letter = worksheet.cell(row=3, column=status_col).column_letter
    for status, (fill_color, font_color) in STATUS_COLORS.items():
        rule = FormulaRule(
            formula=[f'${status_letter}3="{status}"'],
            stopIfTrue=False,
            fill=PatternFill(fill_type="solid", fgColor=fill_color),
            font=Font(color=font_color, bold=True),
        )
        worksheet.conditional_formatting.add(f"{status_letter}3:{status_letter}1048576", rule)


def _apply_common_sheet_polish(worksheet) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A3"
    worksheet.auto_filter.ref = f"A2:{worksheet.cell(row=worksheet.max_row, column=worksheet.max_column).coordinate}"
    worksheet.sheet_properties.tabColor = SHEET_ACCENT_COLOR
    worksheet.row_dimensions[1].height = 24
    worksheet.row_dimensions[2].height = 28

    for cell in worksheet[1]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    for cell in worksheet[2]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for row_index in range(3, worksheet.max_row + 1):
        worksheet.row_dimensions[row_index].height = 24
        for cell in worksheet[row_index]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER


def _auto_fit_columns(worksheet) -> None:
    for column_index, column_cells in enumerate(worksheet.columns, start=1):
        max_length = 0
        column_letter = get_column_letter(column_index)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_length:
                max_length = len(value)
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 60)


def _prepare_export_row(row: dict[str, Any]) -> dict[str, Any]:
    prepared = dict(row)
    if not str(prepared.get("Lead Status", "")).strip():
        prepared["Lead Status"] = "Pending"
    return prepared


def _write_sheet_banner(worksheet, message: str) -> None:
    if worksheet.max_column <= 0:
        return
    existing_value = worksheet.cell(row=1, column=1).value
    if isinstance(existing_value, str) and existing_value.strip() == message:
        return
    worksheet.insert_rows(1)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=worksheet.max_column)
    banner_cell = worksheet.cell(row=1, column=1, value=message)
    banner_cell.fill = PatternFill(fill_type="solid", fgColor=ENGAGEMENT_BAND_FILL)
    banner_cell.font = Font(color=ENGAGEMENT_BAND_TEXT, bold=True, size=11)
    banner_cell.alignment = Alignment(horizontal="left", vertical="center")
    banner_cell.border = THIN_BORDER


@dataclass
class ExportResult:
    total_saved: int
    duplicates_removed: int
    output_path: str


def _load_existing_rows(output_path: str) -> list[dict[str, Any]]:
    if not os.path.exists(output_path):
        return []
    try:
        dataframe = pd.read_excel(output_path, sheet_name=LEADS_SHEET_NAME)
        return dataframe.fillna("").to_dict(orient="records")
    except Exception:
        return []


def _style_leads_sheet(workbook) -> None:
    worksheet = workbook[LEADS_SHEET_NAME]
    header_fill = PatternFill(fill_type="solid", fgColor=EXCEL_HEADER_COLOR)
    header_font = Font(color=EXCEL_HEADER_FONT_COLOR, bold=True)
    alt_fill = PatternFill(fill_type="solid", fgColor=EXCEL_ROW_ALT_COLOR)
    base_fill = PatternFill(fill_type="solid", fgColor=EXCEL_ROW_BASE_COLOR)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row_index in range(2, worksheet.max_row + 1):
        row_fill = alt_fill if row_index % 2 == 1 else base_fill
        for cell in worksheet[row_index]:
            cell.fill = row_fill
    _write_sheet_banner(worksheet, "Radz Scraper lead tracker | Review each lead and update the status column.")
    _apply_common_sheet_polish(worksheet)

    header_index = {cell.value: idx + 1 for idx, cell in enumerate(worksheet[2])}
    website_col = header_index.get("Website")
    maps_col = header_index.get("Google Maps URL")
    rating_col = header_index.get("Rating")
    reviews_col = header_index.get("Total Reviews")

    if website_col:
        for row_index in range(3, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=website_col)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    if maps_col:
        for row_index in range(3, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=maps_col)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    if rating_col:
        for row_index in range(3, worksheet.max_row + 1):
            worksheet.cell(row=row_index, column=rating_col).number_format = "0.0"

    if reviews_col:
        for row_index in range(3, worksheet.max_row + 1):
            worksheet.cell(row=row_index, column=reviews_col).number_format = "#,##0"

    _add_status_validation(worksheet, header_index)
    _auto_fit_columns(worksheet)


def _normalize_export_value(column: str, value: Any) -> Any:
    if pd.isna(value):
        return ""
    if column == "Lead Status":
        return str(value).strip() or "Pending"
    if isinstance(value, str):
        return " ".join(value.split())
    if column == "Rating":
        return float(value) if value != "" else ""
    if column == "Total Reviews":
        try:
            return int(value)
        except Exception:
            return ""
    return value


def _build_upload_ready_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized_rows: list[dict[str, Any]] = []
    for row in rows:
        normalized_rows.append(
            {
                column: _normalize_export_value(column, row.get(column, ""))
                for column in OUTPUT_COLUMNS
            }
        )
    return normalized_rows


def _write_upload_ready_sheet(workbook, rows: list[dict[str, Any]]) -> None:
    if UPLOAD_READY_SHEET_NAME in workbook.sheetnames:
        del workbook[UPLOAD_READY_SHEET_NAME]
    worksheet = workbook.create_sheet(UPLOAD_READY_SHEET_NAME)
    worksheet.append(OUTPUT_COLUMNS)
    for row in rows:
        worksheet.append([row.get(column, "") for column in OUTPUT_COLUMNS])
    _write_sheet_banner(worksheet, "Upload Ready | Clean export for sorting, CRM import, or team review.")

    for cell in worksheet[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor=EXCEL_HEADER_COLOR)
        cell.font = Font(color=EXCEL_HEADER_FONT_COLOR, bold=True)
    _apply_common_sheet_polish(worksheet)
    header_index = {cell.value: idx + 1 for idx, cell in enumerate(worksheet[2])}
    _add_status_validation(worksheet, header_index)
    _auto_fit_columns(worksheet)


def _write_summary_sheet(workbook, metadata: dict[str, Any]) -> None:
    if SUMMARY_SHEET_NAME in workbook.sheetnames:
        del workbook[SUMMARY_SHEET_NAME]
    sheet = workbook.create_sheet(SUMMARY_SHEET_NAME)
    sheet["A1"] = "Radz Scraper summary"
    sheet["A1"].font = Font(bold=True, size=12, color="0F172A")
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor=ENGAGEMENT_BAND_FILL)
    sheet["A1"].border = THIN_BORDER
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.merge_cells("A1:B1")
    rows = [
        ("Search Query", metadata.get("query", "")),
        ("Location", metadata.get("location", "")),
        ("Date & Time Scraped", metadata.get("scraped_at", now_string())),
        ("Total Results Found", metadata.get("total_results_found", 0)),
        ("Lead Rows", metadata.get("lead_rows", 0)),
        ("Social Search Rows", metadata.get("social_rows", 0)),
        ("Universal Finder Rows", metadata.get("universal_rows", 0)),
        ("Website / Directory Rows", metadata.get("website_rows", 0)),
        ("Total Filtered Out", metadata.get("total_filtered_out", 0)),
        ("Total Duplicates Removed", metadata.get("duplicates_removed", 0)),
        ("Total Errors", metadata.get("total_errors", 0)),
        ("Script Version", VERSION),
    ]
    for index, (label, value) in enumerate(rows, start=2):
        label_cell = sheet.cell(row=index, column=1, value=label)
        value_cell = sheet.cell(row=index, column=2, value=value)
        label_cell.font = Font(bold=True, color="0B3558")
        label_cell.fill = PatternFill(fill_type="solid", fgColor=SUMMARY_LABEL_FILL)
        value_cell.fill = PatternFill(fill_type="solid", fgColor=SUMMARY_VALUE_FILL)
        label_cell.border = THIN_BORDER
        value_cell.border = THIN_BORDER
        label_cell.alignment = Alignment(vertical="center")
        value_cell.alignment = Alignment(vertical="center")
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = SHEET_ACCENT_COLOR
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 40


def _write_generic_sheet(workbook, sheet_name: str, columns: list[str], rows: list[dict[str, Any]]) -> None:
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name)
    worksheet.append(columns)
    for row in rows:
        worksheet.append([row.get(column, "") for column in columns])
    _write_sheet_banner(worksheet, f"{sheet_name} | Generated by Radz Scraper for all Social Media.")

    header_fill = PatternFill(fill_type="solid", fgColor=EXCEL_HEADER_COLOR)
    header_font = Font(color=EXCEL_HEADER_FONT_COLOR, bold=True)
    alt_fill = PatternFill(fill_type="solid", fgColor=EXCEL_ROW_ALT_COLOR)
    base_fill = PatternFill(fill_type="solid", fgColor=EXCEL_ROW_BASE_COLOR)

    for cell in worksheet[2]:
        cell.fill = header_fill
        cell.font = header_font

    for row_index in range(3, worksheet.max_row + 1):
        row_fill = alt_fill if row_index % 2 == 1 else base_fill
        for cell in worksheet[row_index]:
            cell.fill = row_fill
    _apply_common_sheet_polish(worksheet)
    _auto_fit_columns(worksheet)


def export_leads(
    rows: list[dict[str, Any]],
    output_path: str,
    append: bool,
    metadata: dict[str, Any],
    logger,
) -> ExportResult:
    ensure_parent_dir(output_path)
    existing_rows = _load_existing_rows(output_path) if append else []
    combined_rows = [_prepare_export_row(row) for row in (existing_rows + rows)]
    deduped_rows, duplicates_removed = deduplicate_leads(combined_rows)
    upload_ready_rows = _build_upload_ready_rows(deduped_rows)

    leads_df = pd.DataFrame(deduped_rows, columns=OUTPUT_COLUMNS)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        leads_df.to_excel(writer, index=False, sheet_name=LEADS_SHEET_NAME)
        workbook = writer.book
        _style_leads_sheet(workbook)
        _write_upload_ready_sheet(workbook, upload_ready_rows)
        summary_metadata = dict(metadata)
        summary_metadata["duplicates_removed"] = duplicates_removed
        _write_summary_sheet(workbook, summary_metadata)

    workbook = load_workbook(output_path)
    _style_leads_sheet(workbook)
    _write_upload_ready_sheet(workbook, upload_ready_rows)
    _write_summary_sheet(workbook, {**metadata, "duplicates_removed": duplicates_removed})
    workbook.save(output_path)

    logger.info("Saved %s deduplicated leads to %s", len(deduped_rows), output_path)
    return ExportResult(
        total_saved=len(deduped_rows),
        duplicates_removed=duplicates_removed,
        output_path=output_path,
    )


def export_workbook(
    rows: list[dict[str, Any]],
    output_path: str,
    append: bool,
    metadata: dict[str, Any],
    logger,
    social_rows: list[dict[str, Any]] | None = None,
    universal_rows: list[dict[str, Any]] | None = None,
) -> ExportResult:
    result = export_leads(rows, output_path, append, metadata, logger)
    if social_rows:
        workbook = load_workbook(output_path)
        _write_generic_sheet(workbook, SOCIAL_SEARCH_SHEET_NAME, SOCIAL_SEARCH_COLUMNS, social_rows)
        if universal_rows:
            _write_generic_sheet(workbook, UNIVERSAL_FINDER_SHEET_NAME, UNIVERSAL_FINDER_COLUMNS, universal_rows)
        _write_summary_sheet(workbook, {**metadata, "duplicates_removed": result.duplicates_removed})
        workbook.save(output_path)
        return ExportResult(
            total_saved=len(universal_rows) if universal_rows else (len(social_rows) if not rows else result.total_saved),
            duplicates_removed=result.duplicates_removed,
            output_path=result.output_path,
        )
    if universal_rows:
        workbook = load_workbook(output_path)
        _write_generic_sheet(workbook, UNIVERSAL_FINDER_SHEET_NAME, UNIVERSAL_FINDER_COLUMNS, universal_rows)
        _write_summary_sheet(workbook, {**metadata, "duplicates_removed": result.duplicates_removed})
        workbook.save(output_path)
        return ExportResult(
            total_saved=len(universal_rows),
            duplicates_removed=result.duplicates_removed,
            output_path=result.output_path,
        )
    return result
